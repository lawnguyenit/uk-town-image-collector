"""CSV, XLSX, evidence-index, and QA-report writer.

All string cells starting with ``=``, ``+``, ``-``, or ``@`` (after leading whitespace) are
prefixed with an apostrophe and stored as explicit text cells in XLSX so Excel never
evaluates source content as a formula.
"""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import RESULT_FIELDS, REVIEW_FIELDS
from .review import ready_record, safe_text, write_csv, write_review_template


DATA_DICTIONARY = {
    "input_id": "Stable identifier supplied in input.csv.",
    "place_name": "Place label supplied in input.csv.",
    "wikipedia_url": "Article URL supplied in input.csv.",
    "canonical_title": "Canonical article title from the Wikipedia HTML response.",
    "canonical_page_url": "Canonical article URL from HTML <link rel='canonical'> or final allowed page URL.",
    "page_id": "Wikipedia article ID (wgArticleId or Action API pageid).",
    "revision_id": "Wikipedia revision ID (wgRevisionId or Action API lastrevid).",
    "selection_method": "infobox_first_photo, pageimages_fallback, or html_pageimage_fallback.",
    "file_title": "Selected Wikimedia Commons file title (without the File: prefix).",
    "caption": "Plain-text caption from the Wikipedia infobox (styles/scripts removed).",
    "commons_file_page_url": "Wikimedia Commons file-description page URL.",
    "original_file_url": "Direct original file URL on upload.wikimedia.org (never downloaded in v1).",
    "thumbnail_url": "Preview thumbnail URL (distinct from original_file_url).",
    "mime_type": "Original file MIME type (for example image/jpeg).",
    "width_px": "Original image width in pixels.",
    "height_px": "Original image height in pixels.",
    "creator_text": "Author/creator text from the file description or explicit copyright notice; never inferred from uploader.",
    "creator_url": "HTTPS link to creator profile or user page when present.",
    "credit_text": "Credit line from the file description when present.",
    "license_short_name": "Selected supported license label (for example CC BY-SA 3.0, CC0 1.0, Public domain).",
    "license_url": "HTTPS license deed or file licensing section URL.",
    "license_evidence": "JSON array preserving every observed license choice, its basis, and which choice was selected.",
    "license_status": "SUPPORTED_METADATA when at least one license satisfies v1 rules; REVIEW_REQUIRED otherwise. Not legal advice.",
    "attribution_text": "Suggested plain-text attribution: file_title — creator_text — license_short_name — commons_file_page_url.",
    "status": "CANDIDATE, NO_IMAGE, NEEDS_REVIEW, INVALID_INPUT, FETCH_ERROR, RATE_LIMITED, BLOCKED, or NOT_PROCESSED.",
    "reason_code": "Machine-readable selection/status reason and provenance flags.",
    "reviewed": "Human review state: PENDING, APPROVED, or REJECTED.",
    "reviewer_note": "Human verification note recorded during review.",
    "retrieved_at_utc": "UTC timestamp when the row's metadata was retrieved.",
    "evidence_path": "Relative path to the per-row field evidence JSON.",
    "source_revision_or_timestamp": "Last-modified text or revision timestamp from Commons.",
    "reviewer_license_choice": "Optional reviewer override when a file offers multiple licenses.",
    "reviewer": "Identifier of the human reviewer who inspected the candidate.",
    "reviewed_at_utc": "UTC timestamp when the human review decision was recorded.",
    "review_fingerprint": "16-hex SHA-256 fingerprint of (input_id, file_title, method, original_url, creator, license); invalidates stale approvals if the file changes.",
}


def write_outputs(run_dir: Path, records: list[dict], run_report: dict) -> list[dict]:
    run_dir.mkdir(parents=True, exist_ok=True)
    results = [{field: record.get(field, "") for field in RESULT_FIELDS} for record in records]
    ready_rows = [candidate for record in records if (candidate := ready_record(record)) is not None]
    ready = [{field: row.get(field, "") for field in RESULT_FIELDS} for row in ready_rows]

    write_csv(results, run_dir / "results.csv", RESULT_FIELDS)
    write_review_template(records, run_dir / "review.csv")
    write_csv(ready, run_dir / "ready_for_import.csv", RESULT_FIELDS)
    _write_evidence_index(run_dir, records)
    _write_workbook(run_dir / "results.xlsx", results, records, ready, run_report)
    _verify_workbook(run_dir / "results.xlsx", len(records), len(ready))

    run_report["artifacts"] = {
        name: {"path": name, "sha256": _file_sha256(run_dir / name)}
        for name in ("results.csv", "results.xlsx", "review.csv", "ready_for_import.csv", "evidence-index.csv")
    }
    (run_dir / "run-report.json").write_text(json.dumps(run_report, ensure_ascii=False, indent=2), encoding="utf-8")
    _write_qa_report(run_dir / "qa-report.md", records, ready, run_report)
    return ready


def _write_workbook(path: Path, results: list[dict], reviews: list[dict], ready: list[dict], report: dict) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "Results", results, RESULT_FIELDS)
    _sheet(workbook, "Review", reviews, REVIEW_FIELDS)
    _sheet(workbook, "Ready for Import", ready, RESULT_FIELDS)
    summary_rows = [
        {"key": key, "value": json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)}
        for key, value in report.items() if key != "request_log"
    ]
    _sheet(workbook, "Run Summary", summary_rows, ["key", "value"])
    dictionary_rows = [{"field": field, "meaning": DATA_DICTIONARY.get(field, "")} for field in REVIEW_FIELDS]
    _sheet(workbook, "Data Dictionary", dictionary_rows, ["field", "meaning"])
    workbook.save(path)


def _sheet(workbook: Workbook, name: str, rows: list[dict], fields: list[str]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(max(1, len(fields)))}{max(1, len(rows) + 1)}"
    sheet.append(fields)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="top", wrap_text=False)
    for row in rows:
        sheet.append([safe_text(row.get(field, "")) for field in fields])
    for col_index, field in enumerate(fields, start=1):
        sample_len = max([len(field)] + [min(len(str(r.get(field, ""))), 48) for r in rows[:25]])
        sheet.column_dimensions[get_column_letter(col_index)].width = min(max(sample_len + 3, 12), 44)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"


def _write_evidence_index(run_dir: Path, records: list[dict]) -> None:
    fields = [
        "input_id", "status", "reason_code", "evidence_path", "evidence_sha256",
        "wikipedia_url", "article_response_sha256", "commons_file_page_url",
        "commons_response_sha256", "retrieved_at_utc",
    ]
    with (run_dir / "evidence-index.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in records:
            ev_rel = row.get("evidence_path", "")
            ev_file = run_dir / ev_rel if ev_rel else None
            ev_hash = _file_sha256(ev_file) if ev_file and ev_file.exists() else ""
            ev = row.get("_evidence") or {}
            writer.writerow({
                "input_id": safe_text(row.get("input_id", "")),
                "status": safe_text(row.get("status", "")),
                "reason_code": safe_text(row.get("reason_code", "")),
                "evidence_path": safe_text(ev_rel),
                "evidence_sha256": ev_hash,
                "wikipedia_url": safe_text(row.get("canonical_page_url") or row.get("wikipedia_url", "")),
                "article_response_sha256": _first_hash(ev.get("article")),
                "commons_file_page_url": safe_text(row.get("commons_file_page_url", "")),
                "commons_response_sha256": _first_hash(ev.get("commons")),
                "retrieved_at_utc": safe_text(row.get("retrieved_at_utc", "")),
            })


def _first_hash(items) -> str:
    if not isinstance(items, list):
        return ""
    for item in items:
        if isinstance(item, dict) and item.get("response_sha256"):
            return str(item["response_sha256"])
    return ""


def _file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_qa_report(path: Path, records: list[dict], ready: list[dict], report: dict) -> None:
    status_counts = report.get("status_counts", {})
    total = len(records)
    candidates = status_counts.get("CANDIDATE", 0)
    needs_review = status_counts.get("NEEDS_REVIEW", 0)
    errors = sum(status_counts.get(k, 0) for k in ("FETCH_ERROR", "INVALID_INPUT", "RATE_LIMITED", "BLOCKED", "NOT_PROCESSED"))
    lines = [
        "# QA Report — UK Town Image & Attribution Collector",
        "",
        f"- **Run ID:** `{report.get('run_id', '')}`",
        f"- **Run status:** `{report.get('run_status', 'UNKNOWN')}`",
        f"- **Started / finished (UTC):** `{report.get('started_at_utc', '')}` → `{report.get('finished_at_utc', '')}`",
        f"- **Input SHA-256:** `{report.get('input_file_sha256', '')}`",
        "",
        "## Separate operational metrics (Brief §10)",
        "",
        "These metrics have different denominators and meanings and are intentionally never merged into a single 'accuracy' number:",
        "",
        "| Metric | Count | Notes |",
        "| --- | --- | --- |",
        f"| Supplied inputs | {total} | Exact row order and unique `input_id` preserved |",
        f"| Unique canonical Wikipedia pages | {report.get('unique_canonical_pages', 0)} | Deduplicated in fetch cache |",
        f"| Unique Wikimedia Commons files | {report.get('unique_image_files', 0)} | At most one candidate per input |",
        f"| Extraction coverage (`CANDIDATE`) | {candidates}/{total} | Candidate found with complete required metadata |",
        f"| Metadata completeness (`SUPPORTED_METADATA`) | {report.get('metadata_complete_count', 0)}/{total} | Creator + supported license + dimensions + URLs + evidence |",
        f"| Rows flagged `NEEDS_REVIEW` | {needs_review}/{total} | Ambiguous, missing, or unsupported source metadata |",
        f"| Error / unprocessed rows | {errors}/{total} | `FETCH_ERROR` / `INVALID_INPUT` / `RATE_LIMITED` / `BLOCKED` / `NOT_PROCESSED` |",
        f"| Human review `APPROVED` | {report.get('review_approved_count', 0)}/{total} | Bound to `(input_id, file_title)` and metadata fingerprint |",
        f"| Human review `REJECTED` | {report.get('review_rejected_count', 0)}/{total} | Explicitly excluded from import |",
        f"| Human review `PENDING` | {report.get('review_pending_count', 0)}/{total} | Awaiting human verification |",
        f"| `ready_for_import.csv` rows | {len(ready)}/{total} | Requires `CANDIDATE` + `APPROVED` + `SUPPORTED_METADATA` + all required fields |",
        f"| HTTP requests / cache hits / retries | {report.get('request_count', 0)} / {report.get('cache_hits', 0)} / {report.get('retry_count', 0)} | Budget cap: {report.get('max_request_budget', 0)} |",
        "",
        "## Status breakdown",
        "",
    ]
    lines.extend([f"- `{status}`: {count}" for status, count in status_counts.items()])
    lines.extend([
        "",
        "## Provenance and review notes",
        "",
        "- `SUPPORTED_METADATA` indicates that the source page states a license in the v1 supported set (`CC BY 2.0/3.0/4.0`, `CC BY-SA 2.0/3.0/4.0`, `CC0 1.0`, or `Public domain` with an explicit basis such as `PD-self`). It is not legal advice or clearance of personality, trademark, or local property rights.",
        "- When a file is multi-licensed (for example `CC BY-SA 3.0` + `GFDL`), `license_evidence` preserves every option and the collector selects the deterministic supported license while flagging `MULTI_LICENSE_SUPPORTED_SELECTED` in `reason_code`; a reviewer can still override via `reviewer_license_choice`.",
        "- When creator attribution comes from an explicit copyright notice on the file page rather than a standard `{{Information}}` table, `reason_code` includes `CREATOR_FROM_COPYRIGHT_NOTICE` so the reviewer can double-check it.",
        "- All CSV/XLSX text cells starting with `=`, `+`, `-`, or `@` are apostrophe-prefixed and typed as strings (`data_type='s'`).",
    ])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _verify_workbook(path: Path, expected_rows: int, expected_ready: int) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        expected_sheets = ["Results", "Review", "Ready for Import", "Run Summary", "Data Dictionary"]
        if workbook.sheetnames != expected_sheets:
            raise RuntimeError(f"Workbook sheets differ from contract: {workbook.sheetnames}")
        if workbook["Results"].max_row != expected_rows + 1:
            raise RuntimeError("Workbook Results row count mismatch")
        if workbook["Ready for Import"].max_row != expected_ready + 1:
            raise RuntimeError("Workbook Ready for Import row count mismatch")
        for sheet_name in expected_sheets:
            for row in workbook[sheet_name].iter_rows(min_row=2):
                for cell in row:
                    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                        raise RuntimeError(f"Unsafe formula cell in {sheet_name}!{cell.coordinate}")
    finally:
        workbook.close()
