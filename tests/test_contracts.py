"""Comprehensive contract and integration tests covering every Brief §10 scenario.

All tests are synthetic and run offline against local fixtures or an in-memory mock
``requests.Session`` — they never contact Wikimedia.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import tempfile
import unittest
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

import requests
from openpyxl import load_workbook

from uk_town_image_collector.exporter import _write_workbook
from uk_town_image_collector.mediawiki_api import clean_extmetadata_html, parse_commons_imageinfo, parse_pageimages_response
from uk_town_image_collector.models import RESULT_FIELDS, REVIEW_FIELDS, FetchResponse
from uk_town_image_collector.parsers import license_status, parse_article, parse_commons, select_license
from uk_town_image_collector.review import apply_review_decisions, load_review_decisions, ready_record, review_fingerprint, write_csv
from uk_town_image_collector.runner import RunInterrupted, run_job
from uk_town_image_collector.transport import FetchError, HtmlClient, StopRun
from uk_town_image_collector.validation import InputValidationError, read_and_validate_input


ROOT = Path(__file__).parent
FIXTURES = ROOT / "fixtures"
ROBOTS_ALLOW = b"User-agent: *\nDisallow: /wiki/Special:\n"


def fixture_response(filename: str, url: str) -> FetchResponse:
    body = (FIXTURES / filename).read_bytes()
    return FetchResponse(url, url, 200, body, hashlib.sha256(body).hexdigest(), "2026-10-26T00:00:00Z")


def html_response(body: bytes, url: str, status: int = 200) -> FetchResponse:
    return FetchResponse(url, url, status, body, hashlib.sha256(body).hexdigest(), "2026-10-26T00:00:00Z")


class FakeHTTPResponse:
    def __init__(self, status_code: int = 200, body: bytes = b"", headers: dict | None = None):
        self.status_code = status_code
        self._body = body
        self.headers = headers or {}

    def iter_content(self, chunk_size: int = 65536):
        for i in range(0, len(self._body), chunk_size):
            yield self._body[i:i + chunk_size]

    def close(self) -> None:
        pass


class FakeSession:
    """Deterministic in-memory session mapping URLs (or a callable router) to responses."""

    def __init__(self, routes):
        self.routes = routes
        self.calls: list[str] = []

    def get(self, url: str, **_kwargs):
        self.calls.append(url)
        handler = self.routes(url) if callable(self.routes) else self.routes.get(url)
        if isinstance(handler, list):
            handler = handler.pop(0)
        if callable(handler):
            handler = handler(url)
        if isinstance(handler, BaseException):
            raise handler
        if handler is None:
            return FakeHTTPResponse(404, b"not found")
        return handler


def make_client(cache_dir: Path, routes, *, max_requests: int = 62) -> HtmlClient:
    return HtmlClient(
        "https://github.com/lawnguyenit",
        cache_dir,
        min_delay_seconds=0.0,
        max_requests=max_requests,
        session=FakeSession(routes),
        sleep=lambda _s: None,
    )


def write_input_csv(path: Path, rows: list[tuple[str, str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["input_id", "place_name", "wikipedia_url"])
        writer.writerows(rows)


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def sample_article_html(title: str, file_title: str = "Town_view.jpg", caption: str = "Town centre") -> bytes:
    slug = title.replace(" ", "_")
    file_slug = file_title.replace(" ", "_")
    return (
        f'<!doctype html><html><head><link rel="canonical" href="https://en.wikipedia.org/wiki/{slug}">'
        f'<script>RLCONF={{"wgArticleId":101,"wgRevisionId":9001}};</script></head>'
        f'<body><h1 id="firstHeading">{title}</h1><table class="infobox">'
        f'<tr><td class="infobox-image"><a href="https://en.wikipedia.org/wiki/File:{file_slug}">'
        f'<img src="https://thumb.wikimedia.org/wikipedia/commons/thumb/a/aa/{file_slug}/250px-{file_slug}" alt="photo"></a>'
        f'<div class="infobox-caption">{caption}</div></td></tr></table></body></html>'
    ).encode("utf-8")


def sample_commons_html(file_title: str = "Town view.jpg", author: str = "Jane Doe", license_short: str = "CC BY-SA 3.0",
                        license_url: str = "https://creativecommons.org/licenses/by-sa/3.0/") -> bytes:
    slug = file_title.replace(" ", "_")
    return (
        f'<!doctype html><html><body><h1 id="firstHeading"><span class="mw-page-title-main">File:{file_title}</span></h1>'
        f'<div id="file"><a class="internal" href="https://upload.wikimedia.org/wikipedia/commons/a/aa/{slug}">'
        f'<img class="thumbimage" src="https://thumb.wikimedia.org/wikipedia/commons/thumb/a/aa/{slug}/400px-{slug}"></a>'
        f'<div class="fileinfo">1,600 × 1,200 pixels, MIME type: image/jpeg</div></div>'
        f'<table class="fileinfotpl-type-information"><tr><th>Author</th>'
        f'<td><a href="https://commons.wikimedia.org/wiki/User:Jane_Doe">{author}</a></td></tr>'
        f'<tr><th>Credit line</th><td>Own work</td></tr></table>'
        f'<div class="licensetpl"><span class="licensetpl_short">{license_short}</span>'
        f'<span class="licensetpl_link">{license_url}</span></div>'
        f'<div id="footer-info-lastmod">This page was last edited on 1 October 2026.</div></body></html>'
    ).encode("utf-8")


class ContractTests(unittest.TestCase):
    # ------------------------------------------------------------------ 1. Input validation
    def test_input_validation_enforces_20_rows_unique_id_and_enwiki_article_url(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "input.csv"
            write_input_csv(path, [(f"UK{i:03}", f"Town {i}", f"https://en.wikipedia.org/wiki/Town_{i}") for i in range(1, 21)])
            records, digest = read_and_validate_input(path)
            self.assertEqual(20, len(records))
            self.assertEqual(64, len(digest))

            # Duplicate input_id fails preflight.
            bad = [(f"UK{i:03}", f"Town {i}", f"https://en.wikipedia.org/wiki/Town_{i}") for i in range(1, 21)]
            bad[1] = ("UK001", "Duplicate ID", "https://en.wikipedia.org/wiki/Town_2")
            write_input_csv(path, bad)
            with self.assertRaises(InputValidationError):
                read_and_validate_input(path)

    def test_invalid_url_host_credentials_and_namespace_emit_invalid_input_row(self):
        with tempfile.TemporaryDirectory() as temp:
            tmp = Path(temp)
            path = tmp / "input.csv"
            write_input_csv(path, [
                ("UK001", "Wrong host", "https://commons.wikimedia.org/wiki/Bibury"),
                ("UK002", "Credentials", "https://user:pass@en.wikipedia.org/wiki/Bibury"),
                ("UK003", "File namespace", "https://en.wikipedia.org/wiki/File:Arlington_Row_Bibury.jpg"),
                ("UK004", "", "https://en.wikipedia.org/wiki/Bibury"),
                ("UK005", "Bibury", "https://en.wikipedia.org/wiki/Bibury"),
            ])
            with self.assertRaises(InputValidationError):
                read_and_validate_input(path, expected_count=5, allow_row_invalid=False)
            records, _ = read_and_validate_input(path, expected_count=5, allow_row_invalid=True)
            self.assertEqual(["DISALLOWED_HOST", "URL_HAS_CREDENTIALS", "NON_ARTICLE_NAMESPACE", "BLANK_PLACE_NAME", ""], [r.error for r in records])

            routes = {
                "https://en.wikipedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://commons.wikimedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://en.wikipedia.org/wiki/Bibury": FakeHTTPResponse(200, sample_article_html("Bibury")),
                "https://commons.wikimedia.org/wiki/File:Town_view.jpg": FakeHTTPResponse(200, sample_commons_html()),
            }
            out = run_job(path, tmp / "out", tmp / "cache", "https://github.com/lawnguyenit",
                          expected_count=5, allow_row_invalid=True, enable_api=False, client=make_client(tmp / "cache", routes))
            results = read_csv_rows(out / "results.csv")
            self.assertEqual(["INVALID_INPUT", "INVALID_INPUT", "INVALID_INPUT", "INVALID_INPUT", "CANDIDATE"], [r["status"] for r in results])

    # ------------------------------------------------------------------ 2. Article parsing & montage caption
    def test_article_selects_first_photo_after_skipping_map_and_strips_style_tags_from_montage(self):
        parsed = parse_article(fixture_response("article_infobox.html", "https://en.wikipedia.org/wiki/Example_Town"))
        self.assertEqual("Example Town", parsed.canonical_title)
        self.assertEqual("Town photo.jpg", parsed.candidate.file_title)
        self.assertEqual("infobox_first_photo", parsed.candidate.selection_method)
        self.assertEqual("The old bridge over the river", parsed.candidate.caption)
        self.assertEqual("CANDIDATE", parsed.status)

        montage_html = (
            b'<!doctype html><html><body><h1 id="firstHeading">Nailsworth</h1><table class="infobox">'
            b'<tr><td class="infobox-full-data"><style>.mw-parser-output .tmulti{display:flex}</style>'
            b'<div class="thumb tmulti"><div class="tsingle"><div><a href="https://en.wikipedia.org/wiki/File:Station.jpg">'
            b'<img src="//thumb.wikimedia.org/wikipedia/commons/thumb/a/aa/Station.jpg/120px-Station.jpg?utm_source=en.wikipedia.org"></a></div>'
            b'<div>Nailsworth Station</div></div></div></td></tr></table></body></html>'
        )
        montage = parse_article(html_response(montage_html, "https://en.wikipedia.org/wiki/Nailsworth"))
        self.assertEqual("Nailsworth Station", montage.candidate.caption)
        self.assertNotIn("utm_source", montage.candidate.thumbnail_url)

    def test_no_image_and_pageimages_api_fallback_and_og_image_fallback(self):
        empty_article = b'<!doctype html><html><head><link rel="canonical" href="https://en.wikipedia.org/wiki/Quiet_Village"></head><body><h1 id="firstHeading">Quiet Village</h1></body></html>'
        parsed_empty = parse_article(html_response(empty_article, "https://en.wikipedia.org/wiki/Quiet_Village"))
        self.assertEqual("NO_IMAGE", parsed_empty.status)

        og_parsed = parse_article(fixture_response("article_fallback.html", "https://en.wikipedia.org/wiki/Fallback_Town"))
        self.assertEqual("html_pageimage_fallback", og_parsed.candidate.selection_method)
        self.assertEqual("Fallback view.jpg", og_parsed.candidate.file_title)

        local_og = parse_article(fixture_response("article_noncommons_fallback.html", "https://en.wikipedia.org/wiki/Local_Town"))
        self.assertEqual("NEEDS_REVIEW", local_og.status)
        self.assertEqual("NON_COMMONS_OG_IMAGE", local_og.reason_code)

        api_payload = {
            "query": {"pages": [{
                "pageid": 555, "lastrevid": 777, "title": "Quiet Village", "pageimage": "Village_green.jpg",
                "thumbnail": {"source": "https://thumb.wikimedia.org/wikipedia/commons/thumb/1/11/Village_green.jpg/400px-Village_green.jpg?utm_campaign=api"},
            }]}
        }
        page_id, rev_id, cand = parse_pageimages_response(api_payload, html_response(b"{}", "https://en.wikipedia.org/w/api.php"))
        self.assertEqual(("555", "777"), (page_id, rev_id))
        self.assertEqual("pageimages_fallback", cand.selection_method)
        self.assertEqual("Village green.jpg", cand.file_title)
        self.assertNotIn("utm_campaign", cand.thumbnail_url)

    # ------------------------------------------------------------------ 3. Commons metadata, multi-license, PD-self, copyright notice
    def test_commons_multilicense_selects_supported_license_and_keeps_all_evidence(self):
        parsed = parse_commons(fixture_response("commons_multilicense.html", "https://commons.wikimedia.org/wiki/File:Town_photo.jpg"))
        self.assertEqual("PhotoMaker", parsed.creator_text)
        self.assertNotIn("UploadBot", parsed.creator_text)
        self.assertEqual(2, len(parsed.licenses))
        chosen = select_license(parsed.licenses)
        self.assertEqual("CC BY-SA 3.0", chosen["name"])
        self.assertEqual("SUPPORTED_METADATA", license_status(parsed.licenses))

    def test_public_domain_requires_explicit_basis_and_copyright_notice_extracts_author(self):
        pd_without_basis = (
            b'<!doctype html><html><body><h1 id="firstHeading">File:Mystery.jpg</h1>'
            b'<div id="file"><a class="internal" href="https://upload.wikimedia.org/wikipedia/commons/a/aa/Mystery.jpg"></a>'
            b'<div class="fileinfo">800 x 600 pixels, MIME type: image/jpeg</div></div>'
            b'<table><tr><th>Author</th><td>Anon</td></tr></table>'
            b'<table class="licensetpl"><tr><td><span class="licensetpl_short">Public domain</span></td></tr></table></body></html>'
        )
        unverified_pd = parse_commons(html_response(pd_without_basis, "https://commons.wikimedia.org/wiki/File:Mystery.jpg"))
        self.assertIsNone(select_license(unverified_pd.licenses))
        self.assertEqual("REVIEW_REQUIRED", license_status(unverified_pd.licenses))

        pd_self_html = (
            b'<!doctype html><html><body><h1 id="firstHeading">File:PainswickView2.jpg</h1>'
            b'<div id="file"><a class="internal" href="https://upload.wikimedia.org/wikipedia/commons/e/e6/PainswickView2.jpg"></a>'
            b'<div class="fileinfo">2,048 x 1,536 pixels, MIME type: image/jpeg</div></div>'
            b'<div id="mw-imagepage-content"><table><tr><th>Author</th><td>Jongleur100</td></tr></table>'
            b'<table class="licensetpl"><tr><td><span class="licensetpl_short">Public domain</span></td></tr></table>'
            b'<table class="layouttemplate"><tr><td>I, the copyright holder of this work, release this work into the public domain. This applies worldwide.</td></tr></table></div></body></html>'
        )
        verified_pd = parse_commons(html_response(pd_self_html, "https://commons.wikimedia.org/wiki/File:PainswickView2.jpg"))
        selected_pd = select_license(verified_pd.licenses)
        self.assertEqual("Public domain", selected_pd["name"])
        self.assertEqual("PD-self", selected_pd["basis"])
        self.assertEqual("SUPPORTED_METADATA", license_status(verified_pd.licenses))

        notice_html = (
            b'<!doctype html><html><body><h1 id="firstHeading">File:Malmesbury.jpg</h1>'
            b'<div id="file"><a class="internal" href="https://upload.wikimedia.org/wikipedia/commons/4/42/Malmesbury.jpg"></a>'
            b'<div class="fileinfo">972 x 1,298 pixels, MIME type: image/jpeg</div></div>'
            b'<div id="mw-imagepage-content"><table><tr><td>Photograph &#169; <a href="https://en.wikipedia.org/wiki/User:Solipsist">Andrew Dunn</a>, 15 September 2005.</td></tr></table>'
            b'<div class="licensetpl"><span class="licensetpl_short">CC BY-SA 2.0</span><span class="licensetpl_link">https://creativecommons.org/licenses/by-sa/2.0</span></div></div></body></html>'
        )
        notice_meta = parse_commons(html_response(notice_html, "https://commons.wikimedia.org/wiki/File:Malmesbury.jpg"))
        self.assertEqual("Andrew Dunn", notice_meta.creator_text)
        self.assertEqual("https://en.wikipedia.org/wiki/User:Solipsist", notice_meta.creator_url)
        self.assertEqual("copyright_notice", notice_meta.creator_source)

    def test_missing_license_stays_review_required(self):
        parsed = parse_commons(fixture_response("commons_missing_license.html", "https://commons.wikimedia.org/wiki/File:No_License.jpg"))
        self.assertEqual([], parsed.licenses)
        self.assertEqual("REVIEW_REQUIRED", license_status(parsed.licenses))

    def test_cc0_one_point_zero_url_is_supported_and_source_label_is_preserved(self):
        fixture = (
            b'<!doctype html><html><body><h1 id="firstHeading"><span class="mw-page-title-main">CC0.jpg</span></h1>'
            b'<div id="file"><a class="internal" href="https://upload.wikimedia.org/wikipedia/commons/a/aa/CC0.jpg">'
            b'<img class="thumbimage" src="https://thumb.wikimedia.org/wikipedia/commons/thumb/a/aa/CC0.jpg/320px-CC0.jpg"></a>'
            b'<div class="fileinfo">640 x 480 pixels, MIME type: image/jpeg</div></div>'
            b'<table><tr><td>Author</td><td>Maker</td></tr></table>'
            b'<table class="licensetpl"><tr><td><span class="licensetpl_short">CC0</span>'
            b'<span class="licensetpl_link">http://creativecommons.org/publicdomain/zero/1.0/deed.en</span></td></tr></table></body></html>'
        )
        parsed = parse_commons(html_response(fixture, "https://commons.wikimedia.org/wiki/File:CC0.jpg"))
        self.assertEqual("CC0 1.0", parsed.licenses[0]["name"])
        self.assertEqual("CC0", parsed.licenses[0]["source_label"])
        self.assertEqual("SUPPORTED_METADATA", license_status(parsed.licenses))

    # ------------------------------------------------------------------ 4. extmetadata sanitization
    def test_extmetadata_html_strips_scripts_and_non_https_creator_links(self):
        dirty = '<script>alert(1)</script><style>body{display:none}</style><a href="javascript:evil()">Bob</a> <a href="http://insecure.example/u">x</a> <a href="https://example.org/bob">Profile</a>'
        text, link = clean_extmetadata_html(dirty)
        self.assertEqual("Bob x Profile", text)
        self.assertEqual("https://example.org/bob", link)

        payload = {"query": {"pages": [{"title": "File:Clean.jpg", "imageinfo": [{
            "url": "https://upload.wikimedia.org/wikipedia/commons/a/aa/Clean.jpg?utm_source=api",
            "thumburl": "https://thumb.wikimedia.org/wikipedia/commons/thumb/a/aa/Clean.jpg/400px-Clean.jpg",
            "width": 1000, "height": 800, "mime": "image/jpeg", "timestamp": "2026-01-01T00:00:00Z", "user": "UploaderNotAuthor",
            "extmetadata": {
                "Artist": {"value": dirty},
                "Credit": {"value": "<b>Own work</b>"},
                "LicenseShortName": {"value": "CC BY 4.0"},
                "LicenseUrl": {"value": "https://creativecommons.org/licenses/by/4.0/"},
            },
        }]}]}}
        meta = parse_commons_imageinfo(payload, html_response(b"{}", "https://commons.wikimedia.org/w/api.php"))
        self.assertEqual("Bob x Profile", meta.creator_text)
        self.assertNotIn("UploaderNotAuthor", meta.creator_text)
        self.assertEqual("https://upload.wikimedia.org/wikipedia/commons/a/aa/Clean.jpg", meta.original_file_url)

    # ------------------------------------------------------------------ 5. Redirect + duplicate canonical page deduplication
    def test_redirect_and_two_inputs_sharing_canonical_page_fetch_once_and_keep_two_rows(self):
        with tempfile.TemporaryDirectory() as temp:
            tmp = Path(temp)
            path = tmp / "input.csv"
            write_input_csv(path, [
                ("UK001", "Bibury", "https://en.wikipedia.org/wiki/Bibury"),
                ("UK002", "Bibury Village Redirect", "https://en.wikipedia.org/wiki/Bibury_Redirect"),
            ])
            routes = {
                "https://en.wikipedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://commons.wikimedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://en.wikipedia.org/wiki/Bibury": FakeHTTPResponse(200, sample_article_html("Bibury")),
                "https://en.wikipedia.org/wiki/Bibury_Redirect": FakeHTTPResponse(302, b"", {"Location": "/wiki/Bibury"}),
                "https://commons.wikimedia.org/wiki/File:Town_view.jpg": FakeHTTPResponse(200, sample_commons_html()),
            }
            out = run_job(path, tmp / "out", tmp / "cache", "https://github.com/lawnguyenit",
                          expected_count=2, enable_api=False, client=make_client(tmp / "cache", routes))
            report = json.loads((out / "run-report.json").read_text(encoding="utf-8"))
            results = read_csv_rows(out / "results.csv")
            self.assertEqual(2, len(results))
            self.assertEqual(1, report["unique_canonical_pages"])
            self.assertEqual(1, report["unique_image_files"])
            # Commons page was fetched only once across both inputs.
            commons_fetches = [r for r in report["request_log"] if r["purpose"] == "commons_file_page"]
            self.assertEqual(1, len(commons_fetches))

    # ------------------------------------------------------------------ 6. HTTP 404, timeout, 5xx retry, 403/429 stop
    def test_404_and_timeout_and_5xx_retry_and_429_stop_leave_remaining_not_processed(self):
        with tempfile.TemporaryDirectory() as temp:
            tmp = Path(temp)
            path = tmp / "input.csv"
            write_input_csv(path, [
                ("UK001", "Missing Town", "https://en.wikipedia.org/wiki/Missing_Town"),
                ("UK002", "Flaky Town", "https://en.wikipedia.org/wiki/Flaky_Town"),
                ("UK003", "Timeout Town", "https://en.wikipedia.org/wiki/Timeout_Town"),
                ("UK004", "RateLimited Town", "https://en.wikipedia.org/wiki/RateLimited_Town"),
                ("UK005", "Unreached Town", "https://en.wikipedia.org/wiki/Unreached_Town"),
            ])
            routes = {
                "https://en.wikipedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://commons.wikimedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://en.wikipedia.org/wiki/Missing_Town": FakeHTTPResponse(404, b"not found"),
                "https://en.wikipedia.org/wiki/Flaky_Town": [
                    FakeHTTPResponse(503, b"busy"),
                    FakeHTTPResponse(200, sample_article_html("Flaky Town")),
                ],
                "https://commons.wikimedia.org/wiki/File:Town_view.jpg": FakeHTTPResponse(200, sample_commons_html()),
                "https://en.wikipedia.org/wiki/Timeout_Town": [
                    requests.Timeout("timed out"),
                    requests.Timeout("timed out"),
                    requests.Timeout("timed out"),
                ],
                "https://en.wikipedia.org/wiki/RateLimited_Town": FakeHTTPResponse(429, b"slow down", {"Retry-After": "60"}),
                "https://en.wikipedia.org/wiki/Unreached_Town": FakeHTTPResponse(200, sample_article_html("Unreached Town")),
            }
            out = run_job(path, tmp / "out", tmp / "cache", "https://github.com/lawnguyenit",
                          expected_count=5, enable_api=False, client=make_client(tmp / "cache", routes))
            report = json.loads((out / "run-report.json").read_text(encoding="utf-8"))
            rows = {r["input_id"]: r for r in read_csv_rows(out / "results.csv")}
            self.assertEqual("STOPPED", report["run_status"])
            self.assertEqual("RATE_LIMITED", report["stop_reason"])
            self.assertEqual(3, report["retry_count"])  # 1 retry on 503 + 2 retries on timeout
            self.assertEqual("FETCH_ERROR", rows["UK001"]["status"])
            self.assertEqual("HTTP_404", rows["UK001"]["reason_code"])
            self.assertEqual("CANDIDATE", rows["UK002"]["status"])
            self.assertEqual("FETCH_ERROR", rows["UK003"]["status"])
            self.assertEqual("NETWORK_ERROR", rows["UK003"]["reason_code"])
            self.assertEqual("RATE_LIMITED", rows["UK004"]["status"])
            self.assertEqual("NOT_PROCESSED", rows["UK005"]["status"])

    def test_403_access_denial_marks_blocked_and_stops_batch(self):
        with tempfile.TemporaryDirectory() as temp:
            tmp = Path(temp)
            path = tmp / "input.csv"
            write_input_csv(path, [
                ("UK001", "Blocked Town", "https://en.wikipedia.org/wiki/Blocked_Town"),
                ("UK002", "Later Town", "https://en.wikipedia.org/wiki/Later_Town"),
            ])
            routes = {
                "https://en.wikipedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://commons.wikimedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://en.wikipedia.org/wiki/Blocked_Town": FakeHTTPResponse(403, b"forbidden"),
            }
            out = run_job(path, tmp / "out", tmp / "cache", "https://github.com/lawnguyenit",
                          expected_count=2, enable_api=False, client=make_client(tmp / "cache", routes))
            rows = {r["input_id"]: r for r in read_csv_rows(out / "results.csv")}
            self.assertEqual("BLOCKED", rows["UK001"]["status"])
            self.assertEqual("NOT_PROCESSED", rows["UK002"]["status"])

    # ------------------------------------------------------------------ 7. Checkpoint & resume
    def test_interrupted_batch_writes_checkpoint_and_resumes_without_duplicate_rows(self):
        with tempfile.TemporaryDirectory() as temp:
            tmp = Path(temp)
            path = tmp / "input.csv"
            write_input_csv(path, [
                ("UK001", "Town One", "https://en.wikipedia.org/wiki/Town_One"),
                ("UK002", "Town Two", "https://en.wikipedia.org/wiki/Town_Two"),
            ])
            interrupted_routes = {
                "https://en.wikipedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://commons.wikimedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://en.wikipedia.org/wiki/Town_One": FakeHTTPResponse(200, sample_article_html("Town One")),
                "https://commons.wikimedia.org/wiki/File:Town_view.jpg": FakeHTTPResponse(200, sample_commons_html()),
                "https://en.wikipedia.org/wiki/Town_Two": KeyboardInterrupt(),
            }
            with self.assertRaises(RunInterrupted) as ctx:
                run_job(path, tmp / "out", tmp / "cache", "https://github.com/lawnguyenit",
                        expected_count=2, enable_api=False, client=make_client(tmp / "cache", interrupted_routes))
            checkpoint_dir = ctx.exception.working_dir
            self.assertTrue((checkpoint_dir / "checkpoint.json").exists())

            resume_routes = {
                "https://en.wikipedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://commons.wikimedia.org/robots.txt": FakeHTTPResponse(200, ROBOTS_ALLOW),
                "https://en.wikipedia.org/wiki/Town_Two": FakeHTTPResponse(200, sample_article_html("Town Two")),
                "https://commons.wikimedia.org/wiki/File:Town_view.jpg": FakeHTTPResponse(200, sample_commons_html()),
            }
            out = run_job(path, tmp / "out", tmp / "cache", "https://github.com/lawnguyenit",
                          expected_count=2, resume_from=checkpoint_dir, enable_api=False,
                          client=make_client(tmp / "cache", resume_routes))
            report = json.loads((out / "run-report.json").read_text(encoding="utf-8"))
            results = read_csv_rows(out / "results.csv")
            self.assertEqual(["UK001"], report["rows_restored_from_checkpoint"])
            self.assertEqual(["UK001", "UK002"], [r["input_id"] for r in results])
            self.assertEqual(["CANDIDATE", "CANDIDATE"], [r["status"] for r in results])

    # ------------------------------------------------------------------ 8. Review binding & invalidation on file/metadata change
    def test_review_decision_invalidates_when_file_or_metadata_fingerprint_changes(self):
        base = {field: "v" for field in RESULT_FIELDS}
        base.update({
            "input_id": "UK001", "file_title": "Photo A.jpg", "selection_method": "infobox_first_photo",
            "original_file_url": "https://upload.wikimedia.org/wikipedia/commons/a/aa/Photo_A.jpg",
            "creator_text": "Alice", "license_short_name": "CC BY-SA 3.0", "license_url": "https://creativecommons.org/licenses/by-sa/3.0/",
            "license_evidence": json.dumps([{"name": "CC BY-SA 3.0", "url": "https://creativecommons.org/licenses/by-sa/3.0/", "selected": True}]),
            "license_status": "SUPPORTED_METADATA", "status": "CANDIDATE",
        })
        fp = review_fingerprint(base)

        # Matching file and fingerprint -> APPROVED is applied and enters ready_for_import.
        rows = [dict(base)]
        stats = apply_review_decisions(rows, {("UK001", "Photo A.jpg"): {
            "reviewed": "APPROVED", "reviewer_note": "Verified", "reviewer_license_choice": "",
            "reviewer": "owner", "reviewed_at_utc": "2026-10-26T03:00:00Z", "review_fingerprint": fp,
        }})
        self.assertEqual(1, stats["applied"])
        self.assertIsNotNone(ready_record(rows[0]))

        # File changed to Photo B.jpg -> old approval on Photo A.jpg is ignored.
        changed_file = [dict(base, file_title="Photo B.jpg")]
        apply_review_decisions(changed_file, {("UK001", "Photo A.jpg"): {"reviewed": "APPROVED", "reviewer_note": "old", "review_fingerprint": fp}})
        self.assertEqual("PENDING", changed_file[0]["reviewed"])

        # Same file title, but creator changed on Commons -> fingerprint mismatch invalidates approval.
        changed_creator = [dict(base, creator_text="Different Author")]
        stats2 = apply_review_decisions(changed_creator, {("UK001", "Photo A.jpg"): {"reviewed": "APPROVED", "reviewer_note": "old", "review_fingerprint": fp}})
        self.assertEqual(1, stats2["invalidated_changed_metadata"])
        self.assertEqual("PENDING", changed_creator[0]["reviewed"])
        self.assertIn("REVIEW_INVALIDATED", changed_creator[0]["reviewer_note"])

    # ------------------------------------------------------------------ 9. Formula escaping & route guard
    def test_csv_and_xlsx_neutralise_formulas_and_preserve_vietnamese_unicode(self):
        with tempfile.TemporaryDirectory() as temp:
            tmp = Path(temp)
            csv_path = tmp / "data.csv"
            write_csv([{"text": " =SUM(A1:A2)", "vn": "Ảnh thị trấn đã kiểm duyệt"}], csv_path, ["text", "vn"])
            with csv_path.open(encoding="utf-8-sig", newline="") as handle:
                row = next(csv.DictReader(handle))
            self.assertEqual("' =SUM(A1:A2)", row["text"])
            self.assertEqual("Ảnh thị trấn đã kiểm duyệt", row["vn"])

            xlsx_path = tmp / "result.xlsx"
            rec = {field: "" for field in RESULT_FIELDS}
            rec.update({"input_id": "UK001", "place_name": "=1+1", "caption": "+cmd", "reviewer_note": "Đã duyệt — hợp lệ", "status": "CANDIDATE"})
            _write_workbook(xlsx_path, [rec], [dict(rec)], [], {"run_status": "COMPLETED"})
            book = load_workbook(xlsx_path, data_only=False)
            self.assertEqual(["Results", "Review", "Ready for Import", "Run Summary", "Data Dictionary"], book.sheetnames)
            self.assertEqual("'=1+1", book["Results"]["B2"].value)
            self.assertEqual("'+cmd", book["Results"]["J2"].value)
            self.assertEqual("s", book["Results"]["B2"].data_type)
            self.assertEqual("Đã duyệt — hợp lệ", book["Results"]["AB2"].value)
            book.close()

    def test_route_guard_is_checked_before_http_or_cache(self):
        with tempfile.TemporaryDirectory() as temp:
            called = []
            client = HtmlClient("https://github.com/lawnguyenit", Path(temp), session=FakeSession(lambda u: called.append(u)))
            with self.assertRaises(StopRun):
                client.fetch("https://en.wikipedia.org/wiki/Example", route_guard=lambda _: False)
            self.assertEqual([], called)


if __name__ == "__main__":
    unittest.main()
